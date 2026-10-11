"""Filtered, formatted Excel records and PDF analysis reports."""
from io import BytesIO
from datetime import datetime
from zoneinfo import ZoneInfo
from xml.sax.saxutils import escape
import pandas as pd
from services import report_service, rules_service, prescription_service


def export_records(frame, filters):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    book=Workbook();sheet=book.active;sheet.title='Approved records'
    sheet.append(['MDRRMC Echague - Approved accident records'])
    sheet.append(['Generated',datetime.now(ZoneInfo('Asia/Manila')).isoformat()])
    sheet.append(['Filters',', '.join(f'{k}: {v}' for k,v in filters.items() if v) or 'All records'])
    sheet.append(['Record count',len(frame)])
    columns=list(frame.columns);sheet.append(columns)
    for row in report_service.records(frame):
        sheet.append([report_service.safe_cell(row.get(c)) for c in columns])
    for cell in sheet[5]:
        cell.fill=PatternFill('solid',fgColor='063C29');cell.font=Font(color='FFFFFF',bold=True);cell.alignment=Alignment(wrap_text=True)
    sheet.row_dimensions[5].height=32
    sheet.freeze_panes='A6'
    if columns:sheet.auto_filter.ref=f'A5:{get_column_letter(len(columns))}{max(5,sheet.max_row)}'
    for i,c in enumerate(columns,1):sheet.column_dimensions[get_column_letter(i)].width=42 if c in ('damage_injury','review_notes') else 23
    sheet.sheet_properties.pageSetUpPr.fitToPage=True
    sheet.page_setup.orientation='landscape';sheet.page_setup.paperSize=sheet.PAPERSIZE_A3;sheet.page_setup.fitToWidth=1;sheet.page_setup.fitToHeight=0
    sheet.print_title_rows='1:5';sheet.sheet_view.showGridLines=False
    output=BytesIO();book.save(output);return output.getvalue()


def analysis_pdf(frame, filters, revision):
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak, KeepTogether
    from services import overview_service as overview
    from services.report_charts import bars, series_chart, accident_map
    from services.contextual_prescriptions import plain_language
    from arm.pipeline import ARMPipeline
    styles = getSampleStyleSheet()
    styles['Normal'].fontSize = 10
    styles['Normal'].leading = 15
    styles['Title'].textColor = colors.HexColor('#063C29')
    styles.add(ParagraphStyle(name='SmallNote', fontSize=8.5, leading=12, textColor=colors.HexColor('#405447')))
    output = BytesIO()
    doc = SimpleDocTemplate(output, pagesize=A4, rightMargin=42, leftMargin=42, topMargin=38, bottomMargin=42,
                            title='Road Accident Analysis Report', author='MDRRMC Echague')
    def p(text, style='Normal'):
        return Paragraph(escape(plain_language(str(text))), styles[style])
    story = [p('Road Accident Analysis Report', 'Title'),
             p('Echague, Isabela | Municipal Disaster Risk Reduction Management Council', 'SmallNote'), Spacer(1, 10)]
    generated = datetime.now(ZoneInfo('Asia/Manila')).strftime('%d %B %Y, %I:%M %p PHT')
    story += [p('Generated: '+generated, 'SmallNote'),
              p('Filters: '+(', '.join(f'{k.replace("_", " ")}: {v}' for k,v in filters.items() if v) or 'All approved records'), 'SmallNote'),
              Spacer(1, 10),
              p('Counts refer to approved records, not necessarily separate incidents. All charts use the selected filters.', 'SmallNote')]
    def chart(title, drawing, note='', rows=None):
        block = [p(title, 'Heading2'), drawing]
        if rows:
            maximum = max(r['count'] for r in rows)
            winners = [plain_language(str(r['label'])) for r in rows if r['count'] == maximum]
            if maximum:
                names = ' and '.join(winners)
                verb = 'has' if len(winners)==1 else 'tie for'
                message = f'<b>{escape(names)}</b> {verb} the highest count, with <b>{maximum:,}</b> recorded crashes' + (' each.' if len(winners)>1 else '.')
            else: message = 'No records with a known category are available for this chart.'
            block.append(Paragraph(message, styles['Normal']))
        if note: block.append(p(note, 'SmallNote'))
        story.append(KeepTogether(block))
    if frame.empty:
        story.append(p('No approved records match these filters.', 'Heading2'))
    else:
        data = overview.build_overview(frame, {})
        groups = overview.build_barangay_breakdown(frame, {})
        yearly = data['yearly_trend']
        chart('Crash records over the years', series_chart([r['year'] for r in yearly], {'Records':[r['count'] for r in yearly]}, line=True), rows=[dict(label=r['year'],count=r['count']) for r in yearly])
        chart('Crash records by month', series_chart([r['label'][:3] for r in data['month_distribution']], {'Records':[r['count'] for r in data['month_distribution']]}), 'Months combine selected years.', rows=data['month_distribution'])
        story.append(PageBreak())
        chart('Crash records by hour', series_chart([r['label'] for r in data['hourly_distribution']], {'Records':[r['count'] for r in data['hourly_distribution']]}, hourly=True), rows=data['hourly_distribution'])
        story.append(PageBreak())
        chart('Crash records by day of week', bars(data['day_distribution']), rows=data['day_distribution'])
        chart('Leading recorded causes', bars(data['top_causes']), rows=data['top_causes'])
        story.append(PageBreak())
        chart('Collision types', bars(data['collision_types']), rows=data['collision_types'])
        chart('Vehicles involved', bars(data['top_vehicles']), rows=data['top_vehicles'])
        story.append(PageBreak())
        chart('Barangays with the most crash records', bars(data['top_barangays']), rows=data['top_barangays'])
        chart('Top barangays by day of week', series_chart([d[:3] for d in groups['days']], groups['by_day']), rows=[dict(label=f'{name} on {day}',count=count) for name,values in groups['by_day'].items() for day,count in zip(groups['days'],values)])
        story.append(PageBreak())
        chart('Top barangays by month', series_chart(groups['months'], groups['by_month']), rows=[dict(label=f'{name} in {month}',count=count) for name,values in groups['by_month'].items() for month,count in zip(groups['months'],values)])
        story.append(PageBreak())
        map_image, unmatched, peak = accident_map(frame)
        chart('Accident map - Echague barangays', map_image,
              f'Pale orange: fewer records. Deeper red: more records (highest barangay count: {peak}). Unshaded areas have no matched records. Soft edges are a visual treatment, not exact crash locations or a prediction of danger. {unmatched} records could not be matched to a barangay boundary.', rows=data['top_barangays'])
        story.append(p('Boundaries: philippines-json-maps, 2019 edition (MIT license). Barangay totals include records with missing coordinates.', 'SmallNote'))
        candidate = ARMPipeline(); candidate.data = frame.copy(); candidate._staging = True; candidate.run(force=True)
        ranked, _ = rules_service.filter_rules(candidate.rules, limit=5000)
        rules, seen = [], set()
        for rule in ranked:
            key = frozenset(tuple(i) for i in rule['items'])
            if key not in seen:
                rules.append(rule); seen.add(key)
            if len(rules) == 10: break
        story += [PageBreak(), p('Patterns and suggested actions', 'Heading1'),
                  p('These conditions often appear together; this does not prove that one caused another. Up to 10 different patterns are shown. Each appears in at least 5% of selected records; the second condition appears in at least 60% of records with the first, and at least 20% more often than its usual frequency.', 'SmallNote'),
                  p('Engineering actions require a site check and road-owner approval. Select signs, crossings, barriers or speed-control measures to suit the road; the pattern alone does not confirm a defect.', 'SmallNote'), Spacer(1, 10)]
        if not rules: story.append(p('No patterns meet the selection rules for these records.'))
        for index, rule in enumerate(rules, 1):
            heading = f'{index}. {rule["antecedent_label"]} / {rule["consequent_label"]}'
            block = [p(heading, 'Heading3'), p(f'Both appear in {rule["support"]:.1%} of selected records. Among records with {rule["antecedent_label"]}, {rule["confidence"]:.1%} also include {rule["consequent_label"]}.', 'SmallNote')]
            for action in prescription_service.get().match_rule(rule):
                label = {'enforcement':'Barangay Tanod', 'engineering':'Engineering', 'education':'Public Information & Education'}[action['agency']]
                block.extend([p(label, 'Heading4'), p(action['action'])])
            story.append(KeepTogether(block+[Spacer(1, 12)]))
        story += [Spacer(1, 12), p('Putting the actions into practice', 'Heading2'),
                  p('Have the municipal road-safety team select priority sites, name a responsible office and set a target date. Log hazards, patrol hours and repairs. Review speeds, observed unsafe behavior and new crash records after implementation; a short-term change in counts alone does not prove that an action worked.'),
                  p('Guidance used', 'Heading2'),
                  p('WHO Save LIVES road-safety package: speed management, safer roads, enforcement and emergency response. https://www.who.int/publications/i/item/save-lives-a-road-safety-technical-package', 'SmallNote'),
                  p('DILG guidance on barangay tanods: traffic assistance, patrols and reporting, with police and emergency-service coordination. https://region6.dilg.gov.ph/ligtas-na-paaralan-dilg-deploys-barangay-tanods-to-strengthen-school-safety-nationwide/', 'SmallNote')]
    def footer(canvas, document):
        canvas.setFont('Helvetica', 8); canvas.setFillColor(colors.HexColor('#405447'))
        canvas.drawString(42, 24, 'MDRRMC Echague | Road Accident Analysis Report')
        canvas.drawRightString(A4[0]-42, 24, f'Page {document.page}')
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue()
